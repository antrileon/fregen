from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/guion_capcut_ejercicios_regenmove.docx")
TXT_OUT = Path("docs/prompt_capcut_batch_10s_regenmove.txt")
INDIVIDUAL_DIR = Path("docs/capcut_prompts_10s")


EXERCISES = [
    {
        "area": "Cuello",
        "name": "Deslizamiento del nervio mediano",
        "objective": "Liberar tension neural cervical y disminuir dolor que baja hacia hombro, brazo o mano.",
        "visual": "Persona de pie o sentada, torso recto. Brazo abre a 45 grados con palma hacia arriba. Muneca y dedos se extienden suavemente. La cabeza se inclina al lado contrario y vuelve al centro.",
        "voice": [
            "Colocate sentado o de pie con los hombros relajados.",
            "Abre el brazo a cuarenta y cinco grados con la palma hacia arriba.",
            "Extiende suavemente la muneca y los dedos, sin buscar dolor.",
            "Inclina la cabeza hacia el lado contrario y vuelve al centro.",
            "Repite lento. Debes sentir tension ligera, no corriente ni hormigueo fuerte.",
        ],
        "screen": "Tension suave. No forzar. 8 repeticiones.",
        "avoid": "No subir el hombro, no bloquear el codo con fuerza, no mantener hormigueo.",
        "pain": "Detener si aumenta hormigueo, corriente o dolor irradiado.",
    },
    {
        "area": "Cuello",
        "name": "Liberacion suboccipital con respiracion",
        "objective": "Reducir compresion cervical alta y relajar la base del craneo.",
        "visual": "Persona acostada boca arriba. Toalla enrollada bajo la base del craneo. La barbilla hace un movimiento pequeno hacia atras mientras respira.",
        "voice": [
            "Acuestate boca arriba y coloca una toalla enrollada bajo la base del craneo.",
            "Deja que el peso de la cabeza descanse sobre la toalla.",
            "Haz una pequena retraccion de barbilla, como si hicieras una papada suave.",
            "Respira lento y relaja el cuello durante cinco respiraciones.",
            "La presion debe sentirse comoda, nunca intensa.",
        ],
        "screen": "Respira. Relaja cuello. 5 respiraciones.",
        "avoid": "No levantar la cabeza, no presionar demasiado, no aguantar la respiracion.",
        "pain": "Detener si aparece mareo, nausea o presion intensa.",
    },
    {
        "area": "Hombros",
        "name": "Deslizamiento del nervio radial",
        "objective": "Aliviar irritacion neural asociada a dolor lateral de brazo y hombro.",
        "visual": "Persona de pie, brazo al costado. Hombro bajo. El pulgar gira hacia atras, la muneca flexiona suavemente y la cabeza acompana al lado contrario.",
        "voice": [
            "De pie, deja el brazo relajado al costado.",
            "Mantén el hombro bajo y largo.",
            "Gira el brazo hacia adentro, con el pulgar apuntando hacia atras.",
            "Flexiona la muneca suavemente e inclina la cabeza al lado contrario.",
            "Vuelve al centro y repite lento.",
        ],
        "screen": "Hombro bajo. Movimiento suave. 8 repeticiones.",
        "avoid": "No encoger hombros, no provocar dolor electrico, no hacerlo rapido.",
        "pain": "Trabajar solo con tension suave. Detener si hay dolor electrico.",
    },
    {
        "area": "Hombros",
        "name": "Retracciones escapulares",
        "objective": "Reforzar omoplatos para reducir compresion mecanica en cuello y hombro.",
        "visual": "Persona de pie, brazos al costado. Los omoplatos se acercan suavemente hacia atras y abajo. Mantiene tres segundos y relaja.",
        "voice": [
            "Ponte de pie con los brazos relajados.",
            "Lleva los omoplatos suavemente hacia atras, como si quisieras juntarlos.",
            "Mantén tres segundos sin subir los hombros.",
            "Relaja y vuelve a la posicion inicial.",
            "Haz el movimiento controlado, sin arquear la espalda baja.",
        ],
        "screen": "Omoplatos atras y abajo. 10 repeticiones.",
        "avoid": "No subir hombros, no arquear la espalda baja, no apretar con demasiada fuerza.",
        "pain": "Reducir rango si aparece pinchazo en la parte frontal del hombro.",
    },
    {
        "area": "Espalda alta",
        "name": "Apertura toracica con liberacion neural",
        "objective": "Disminuir compresion postural y mejorar movilidad de espalda alta.",
        "visual": "Persona sentada en silla. Respaldo a media espalda. Brazos cruzados en el pecho. Extiende suavemente sobre el respaldo y regresa.",
        "voice": [
            "Sientate en una silla con el respaldo a media espalda.",
            "Cruza los brazos sobre el pecho.",
            "Abre el pecho hacia atras de forma suave.",
            "Respira y vuelve al centro con control.",
            "El movimiento debe venir de la espalda alta, no del cuello.",
        ],
        "screen": "Abre pecho. Respira. 8 repeticiones.",
        "avoid": "No forzar la zona lumbar, no mover solo el cuello, no contener la respiracion.",
        "pain": "Evitar dolor punzante o presion intensa.",
    },
    {
        "area": "Espalda baja",
        "name": "Deslizamiento del nervio ciatico",
        "objective": "Reducir sensacion de compresion lumbar que baja hacia gluteo o pierna.",
        "visual": "Persona sentada al borde de una silla. Espalda alta. Una pierna se extiende al frente, los dedos del pie apuntan hacia la persona y luego la pierna baja.",
        "voice": [
            "Sientate al borde de una silla con la espalda alta.",
            "Estira una rodilla al frente sin bloquear con fuerza.",
            "Lleva los dedos del pie hacia ti.",
            "Luego baja la pierna y relaja el pie.",
            "Repite como un deslizamiento, no como un estiramiento sostenido.",
        ],
        "screen": "Desliza, no estires. 8 por pierna.",
        "avoid": "No mantener la posicion, no redondear la espalda, no buscar dolor fuerte.",
        "pain": "No debe dejar dolor residual en la pierna despues de la serie.",
    },
    {
        "area": "Espalda baja",
        "name": "Puente corto de descarga lumbar",
        "objective": "Reforzar gluteos para quitar carga repetida de la zona lumbar.",
        "visual": "Persona boca arriba, rodillas flexionadas y pies apoyados. Eleva cadera hasta alinear tronco y baja lento.",
        "voice": [
            "Acuestate boca arriba con rodillas flexionadas y pies apoyados.",
            "Activa los gluteos suavemente.",
            "Eleva la cadera hasta formar una linea comoda con el tronco.",
            "Baja lento, vertebra por vertebra.",
            "Mantén cuello y hombros relajados.",
        ],
        "screen": "Activa gluteos. 8 repeticiones.",
        "avoid": "No hiperextender la espalda baja, no empujar con el cuello, no dejar caer rodillas hacia adentro.",
        "pain": "Suspender si aumenta dolor lumbar durante la elevacion.",
    },
    {
        "area": "Cadera",
        "name": "Liberacion piriforme asistida",
        "objective": "Disminuir compresion en gluteo profundo y mejorar rotacion de cadera.",
        "visual": "Persona boca arriba. Tobillo cruzado sobre rodilla contraria. Acerca la pierna hacia el pecho hasta tension suave y respira.",
        "voice": [
            "Acuestate boca arriba.",
            "Cruza un tobillo sobre la rodilla contraria.",
            "Acerca la pierna hacia el pecho solo hasta sentir tension suave en el gluteo.",
            "Respira lento y suelta.",
            "Repite sin rebotes y sin jalar con fuerza.",
        ],
        "screen": "Tension en gluteo. 6 respiraciones por lado.",
        "avoid": "No torcer la rodilla, no elevar el cuello, no jalar fuerte.",
        "pain": "Debe sentirse en gluteo, no como corriente por la pierna.",
    },
    {
        "area": "Rodillas",
        "name": "Deslizamiento femoral suave",
        "objective": "Reducir tension anterior del muslo que altera mecanica de rodilla.",
        "visual": "Persona acostada de lado. Rodillas ligeramente flexionadas. Lleva el talon hacia gluteo hasta tension leve y regresa.",
        "voice": [
            "Acuestate de lado con las rodillas ligeramente flexionadas.",
            "Mantén la pelvis estable.",
            "Lleva el talon hacia el gluteo hasta una tension leve en el frente del muslo.",
            "Regresa lentamente.",
            "Hazlo suave, como un bombeo, no como un estiramiento fuerte.",
        ],
        "screen": "Pelvis estable. 8 por lado.",
        "avoid": "No arquear la espalda, no jalar fuerte el pie, no provocar dolor frontal agudo.",
        "pain": "No debe aumentar dolor de rodilla ni lumbar.",
    },
    {
        "area": "Rodillas",
        "name": "Sentadilla a silla",
        "objective": "Mejorar tolerancia de carga y control de rodilla en actividades diarias.",
        "visual": "Persona de pie frente a una silla. Lleva cadera atras, toca silla sin desplomarse y sube empujando el piso.",
        "voice": [
            "Ponte de pie frente a una silla estable.",
            "Coloca los pies al ancho de la cadera.",
            "Lleva la cadera hacia atras como si fueras a sentarte.",
            "Toca la silla sin dejarte caer.",
            "Sube empujando el piso y manteniendo las rodillas alineadas.",
        ],
        "screen": "Cadera atras. Rodillas alineadas. 8 repeticiones.",
        "avoid": "No dejar caer las rodillas hacia adentro, no desplomarse en la silla, no levantar talones.",
        "pain": "Reducir profundidad si el dolor supera 3 de 10.",
    },
    {
        "area": "Tobillos",
        "name": "Deslizamiento tibial con bombeo de tobillo",
        "objective": "Disminuir rigidez y sensacion de atrapamiento en pie y tobillo.",
        "visual": "Persona sentada con pierna extendida. Apunta los dedos hacia adelante y luego lleva los dedos hacia si. Movimiento continuo y suave.",
        "voice": [
            "Sientate con una pierna extendida.",
            "Apunta los dedos del pie hacia adelante.",
            "Luego lleva los dedos hacia ti.",
            "Repite de forma suave y continua.",
            "El movimiento debe sentirse como movilidad, no como pinchazo.",
        ],
        "screen": "Punta y flexion. 15 repeticiones.",
        "avoid": "No mover toda la pierna, no bloquear la rodilla, no hacerlo con dolor agudo.",
        "pain": "Detener si aparece pinchazo articular o dolor que aumenta.",
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width


def style_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, RGBColor(46, 116, 181), 18, 10),
        ("Heading 2", 13, RGBColor(46, 116, 181), 14, 7),
        ("Heading 3", 12, RGBColor(31, 77, 120), 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_kv_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    set_table_width(table, [Inches(1.5), Inches(5.25)])
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        set_cell_shading(cells[0], "E8EEF5")
        for cell in cells:
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(10)
        for run in cells[0].paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(31, 77, 120)
    return table


def add_numbered_steps(doc, steps):
    for step in steps:
        paragraph = doc.add_paragraph(step, style="List Number")
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.25


def build_doc():
    doc = Document()
    style_doc(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("Guion para videos: ejercicios RegenMove")
    run.font.name = "Calibri"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(17, 24, 39)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle.add_run(
        "Instrucciones listas para producir clips cortos en CapCut IA. "
        "Cada ejercicio esta pensado para un video de 10 a 20 segundos."
    )

    note = doc.add_paragraph()
    note.paragraph_format.space_after = Pt(12)
    note.add_run("Nota de seguridad: ").bold = True
    note.add_run(
        "estos guiones no reemplazan evaluacion medica. El movimiento debe ser suave, "
        "sin dolor electrico, mareo, perdida de fuerza o aumento de sintomas."
    )

    doc.add_heading("Formato sugerido para todos los videos", level=1)
    add_kv_table(
        doc,
        [
            ("Duracion", "10 segundos para demo rapida o 20 segundos si se agrega voz completa."),
            ("Plano", "Cuerpo visible, fondo limpio, ropa deportiva clara, movimiento lento."),
            ("Texto en pantalla", "Nombre del ejercicio + regla principal de seguridad."),
            ("Ritmo", "Inicio 2s, ejecucion 6s, cierre/recordatorio 2s."),
            ("Estilo", "Clinico, limpio, sin dramatizar dolor. Mostrar control, respiracion y suavidad."),
        ],
    )

    doc.add_heading("Prompt maestro para CapCut IA", level=1)
    doc.add_paragraph(
        "Usa este bloque si quieres generar todos los clips con el mismo estilo. "
        "Cada segmento dura 10 segundos y debe mantener la misma camara, fondo, ropa, ritmo visual, "
        "texto en pantalla y tono clinico."
    )
    add_kv_table(
        doc,
        [
            ("Estilo unico", "Video clinico moderno, fondo claro, ropa deportiva neutra, cuerpo visible, luz suave, movimiento lento y preciso."),
            ("Estructura fija", "0-2s: nombre y posicion inicial. 2-8s: ejecucion del movimiento. 8-10s: regla de seguridad."),
            ("Texto fijo", "Nombre del ejercicio, zona, repeticion/respiracion y regla de dolor."),
            ("Voz", "Narracion breve, calmada y directa. No usar lenguaje exagerado ni prometer cura inmediata."),
            ("Salida", "Crear clips separados de 10 segundos, uno por ejercicio, todos con el mismo formato visual."),
        ],
    )

    current_area = None
    for exercise in EXERCISES:
        if exercise["area"] != current_area:
            current_area = exercise["area"]
            doc.add_heading(current_area, level=1)

        doc.add_heading(exercise["name"], level=2)
        add_kv_table(
            doc,
            [
                ("Objetivo", exercise["objective"]),
                ("Visual", exercise["visual"]),
                ("Texto en pantalla", exercise["screen"]),
                ("Evitar", exercise["avoid"]),
                ("Regla de dolor", exercise["pain"]),
            ],
        )
        doc.add_heading("Voz en off / instrucciones", level=3)
        add_numbered_steps(doc, exercise["voice"])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    TXT_OUT.write_text(build_batch_prompt(), encoding="utf-8")
    write_individual_prompts()
    return OUT


def safe_slug(value):
    allowed = []
    for char in value.lower():
        if char.isalnum():
            allowed.append(char)
        elif char in {" ", "-", "_"}:
            allowed.append("_")
    slug = "".join(allowed)
    while "__" in slug:
        slug = slug.replace("__", "_")
    return slug.strip("_")


def shared_style_block():
    return "\n".join(
        [
            "ESTILO FIJO PARA TODOS LOS VIDEOS REGENMOVE:",
            "- Duracion exacta: 10 segundos.",
            "- Video clinico moderno, fondo claro y limpio.",
            "- Una persona con ropa deportiva neutra.",
            "- Cuerpo visible y movimiento lento, controlado y seguro.",
            "- Camara fija, plano medio o cuerpo completo segun ejercicio.",
            "- Luz suave, sin dramatizar dolor.",
            "- Texto en pantalla minimalista.",
            "- Voz en off calmada, profesional y breve.",
            "- No prometer cura inmediata.",
            "",
            "ESTRUCTURA OBLIGATORIA DEL CLIP:",
            "0-2s: mostrar nombre del ejercicio y posicion inicial.",
            "2-8s: mostrar la ejecucion del movimiento con ritmo lento.",
            "8-10s: mostrar regla de seguridad o dolor.",
        ]
    )


def build_single_prompt(index, exercise):
    voice = " ".join(exercise["voice"])
    return "\n".join(
        [
            shared_style_block(),
            "",
            f"VIDEO {index:02d} DE 11",
            f"Nombre del ejercicio: {exercise['name']}",
            f"Zona: {exercise['area']}",
            f"Objetivo: {exercise['objective']}",
            "",
            "INSTRUCCION VISUAL:",
            exercise["visual"],
            "",
            "TEXTO EN PANTALLA:",
            exercise["screen"],
            "",
            "VOZ EN OFF:",
            voice,
            "",
            "REGLA FINAL DE SEGURIDAD:",
            exercise["pain"],
            "",
            "IMPORTANTE:",
            "Generar solo este ejercicio como un clip independiente de 10 segundos.",
            "Mantener este mismo estilo para todos los demas prompts de RegenMove.",
        ]
    )


def write_individual_prompts():
    INDIVIDUAL_DIR.mkdir(parents=True, exist_ok=True)
    for old_file in INDIVIDUAL_DIR.glob("*.txt"):
        old_file.unlink()
    for index, exercise in enumerate(EXERCISES, start=1):
        path = INDIVIDUAL_DIR / f"{index:02d}_{safe_slug(exercise['name'])}.txt"
        path.write_text(build_single_prompt(index, exercise), encoding="utf-8")


def build_batch_prompt():
    lines = [
        "PROMPT MAESTRO PARA CAPCUT IA - REGENMOVE",
        "",
        "Genera una serie de clips separados de 10 segundos, uno por ejercicio.",
        "Todos los clips deben tener exactamente el mismo estilo visual:",
        "- video clinico moderno",
        "- fondo claro y limpio",
        "- una persona con ropa deportiva neutra",
        "- cuerpo visible y movimiento lento",
        "- camara fija, plano medio o cuerpo completo segun ejercicio",
        "- luz suave, sin dramatizar dolor",
        "- texto en pantalla limpio y minimalista",
        "",
        "Estructura de cada clip de 10 segundos:",
        "0-2s: mostrar nombre del ejercicio y posicion inicial.",
        "2-8s: mostrar la ejecucion del movimiento con ritmo lento.",
        "8-10s: mostrar regla de seguridad o dolor.",
        "",
        "Voz en off para todos: tono calmado, profesional, breve. No prometer cura inmediata. Usar palabras como liberar, aliviar, reforzar y mejorar movilidad.",
        "",
        "CLIPS A GENERAR:",
        "",
    ]

    for index, exercise in enumerate(EXERCISES, start=1):
        voice = " ".join(exercise["voice"])
        lines.extend(
            [
                f"CLIP {index:02d} - {exercise['name']}",
                f"Zona: {exercise['area']}",
                f"Objetivo: {exercise['objective']}",
                f"Visual 10s: {exercise['visual']}",
                f"Texto en pantalla: {exercise['screen']}",
                f"Voz en off corta: {voice}",
                f"Regla de seguridad: {exercise['pain']}",
                "",
            ]
        )

    lines.extend(
        [
            "Indicacion final:",
            "Mantener el mismo personaje, fondo, iluminacion, tipografia y estilo en todos los clips.",
            "Exportar cada ejercicio como clip independiente de 10 segundos, pero generado desde este mismo guion maestro.",
        ]
    )
    return "\n".join(lines)


if __name__ == "__main__":
    print(build_doc())
